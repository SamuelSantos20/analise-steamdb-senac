"""Carrega as relações entre jogos, idiomas e plataformas da planilha completa."""

import argparse
import csv
import os
import unicodedata
from pathlib import Path

import mysql.connector
from dotenv import load_dotenv
from openpyxl import load_workbook


RAIZ = Path(__file__).resolve().parents[2]
PLANILHA = RAIZ / "dados" / "tratados" / "planilhas" / "steam_2026_organizada.xlsx"
RELATORIO = RAIZ / "resultados" / "relatorios" / "idiomas_nao_mapeados.csv"


def normalizar(valor):
    return unicodedata.normalize("NFKC", str(valor).strip()).casefold()


def idiomas_da_celula(valor):
    if valor is None or valor == "[]":
        return []
    if not isinstance(valor, str) or not valor.startswith("[") or not valor.endswith("]"):
        return [str(valor)]
    # A fonte contém entradas como K'iche sem aspas de abertura. Dividir pelos
    # separadores da lista preserva o nome e evita executar conteúdo da célula.
    return [parte.strip().strip("'\" ") for parte in valor[1:-1].split(",") if parte.strip("'\" ")]


def ler_catalogos(cursor):
    cursor.execute("SELECT jogo_id FROM jogos")
    jogos = {int(jogo_id) for (jogo_id,) in cursor}
    cursor.execute("SELECT idioma_id, nome FROM idiomas")
    idiomas = {normalizar(nome): int(idioma_id) for idioma_id, nome in cursor}
    cursor.execute("SELECT plataforma_id, nome FROM plataformas")
    plataformas = {normalizar(nome): int(plataforma_id) for plataforma_id, nome in cursor}
    return jogos, idiomas, plataformas


def extrair_relacoes(jogos, idiomas, plataformas):
    planilha = load_workbook(PLANILHA, read_only=True, data_only=True)
    pares_plataformas = set()
    pares_idiomas = set()
    nao_mapeados = []
    try:
        aba = planilha["Jogos"]
        linhas = aba.iter_rows(values_only=True)
        colunas = {nome: indice for indice, nome in enumerate(next(linhas))}
        ids_jogos = set()
        for linha in linhas:
            jogo_id = int(linha[colunas["AppID"]])
            if jogo_id in ids_jogos or jogo_id not in jogos:
                raise ValueError(f"AppID duplicado ou ausente no banco: {jogo_id}")
            ids_jogos.add(jogo_id)
            for nome in ("Windows", "Mac", "Linux"):
                valor = normalizar(linha[colunas[nome]])
                if valor == "sim":
                    plataforma_id = plataformas.get(normalizar(nome))
                    if plataforma_id is None:
                        raise ValueError(f"Plataforma ausente no catálogo: {nome}")
                    pares_plataformas.add((jogo_id, plataforma_id))
                elif valor != "não":
                    raise ValueError(f"Indicador de plataforma inesperado: {jogo_id}, {nome}, {valor}")
        if ids_jogos != jogos:
            raise ValueError(f"A aba Jogos não cobre {len(jogos - ids_jogos)} jogos do banco")

        aba = planilha["Detalhes e links"]
        linhas = aba.iter_rows(values_only=True)
        colunas = {nome: indice for indice, nome in enumerate(next(linhas))}
        ids_detalhes = set()
        for linha in linhas:
            jogo_id = int(linha[colunas["AppID"]])
            if jogo_id in ids_detalhes or jogo_id not in jogos:
                raise ValueError(f"AppID duplicado ou ausente no banco: {jogo_id}")
            ids_detalhes.add(jogo_id)
            for nome in idiomas_da_celula(linha[colunas["Idiomas disponíveis"]]):
                idioma_id = idiomas.get(normalizar(nome))
                if idioma_id is None:
                    nao_mapeados.append((jogo_id, nome))
                else:
                    pares_idiomas.add((jogo_id, idioma_id))
        if ids_detalhes != jogos:
            raise ValueError(f"A aba Detalhes e links não cobre {len(jogos - ids_detalhes)} jogos do banco")
    finally:
        planilha.close()
    return pares_idiomas, pares_plataformas, nao_mapeados


def inserir_em_lotes(cursor, conexao, tabela, coluna, pares):
    comando = (
        f"INSERT INTO {tabela} (jogo_id, {coluna}) VALUES (%s, %s) "
        f"ON DUPLICATE KEY UPDATE {coluna} = VALUES({coluna})"
    )
    lote = []
    for par in pares:
        lote.append(par)
        if len(lote) == 2000:
            cursor.executemany(comando, lote)
            conexao.commit()
            lote.clear()
    if lote:
        cursor.executemany(comando, lote)
        conexao.commit()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--apply", action="store_true", help="Insere as relações após a validação")
    args = parser.parse_args()
    load_dotenv(RAIZ / ".env")
    conexao = mysql.connector.connect(
        host=os.environ["MYSQL_HOST"],
        user=os.environ["MYSQL_USER"],
        password=os.environ["MYSQL_PASSWORD"],
        database=os.environ["MYSQL_DATABASE"],
    )
    try:
        cursor = conexao.cursor()
        jogos, idiomas, plataformas = ler_catalogos(cursor)
        pares_idiomas, pares_plataformas, nao_mapeados = extrair_relacoes(
            jogos, idiomas, plataformas
        )
        print(f"Jogos validados: {len(jogos)}")
        print(f"Relações de idiomas: {len(pares_idiomas)}")
        print(f"Relações de plataformas: {len(pares_plataformas)}")
        print(f"Nomes de idiomas sem correspondência: {len(nao_mapeados)}")

        RELATORIO.parent.mkdir(parents=True, exist_ok=True)
        with RELATORIO.open("w", encoding="utf-8-sig", newline="") as arquivo:
            escritor = csv.writer(arquivo, delimiter=";")
            escritor.writerow(("jogo_id", "idioma_origem"))
            escritor.writerows(nao_mapeados)
        print(f"Relatório: {RELATORIO}")

        if args.apply:
            inserir_em_lotes(cursor, conexao, "jogos_idiomas", "idioma_id", pares_idiomas)
            inserir_em_lotes(cursor, conexao, "jogos_plataformas", "plataforma_id", pares_plataformas)
            for tabela in ("jogos_idiomas", "jogos_plataformas"):
                cursor.execute(f"SELECT COUNT(*) FROM {tabela}")
                print(f"{tabela}: {cursor.fetchone()[0]} linhas no banco")
        cursor.close()
    finally:
        conexao.close()


if __name__ == "__main__":
    main()
